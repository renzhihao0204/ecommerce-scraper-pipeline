import csv
import json
import os
from pathlib import Path
from typing import List
from models import ProductItem

def export_to_json(items: List[ProductItem], filepath: Path) -> None:
    data = [item.model_dump(mode="json") for item in items]
    filepath.parent.mkdir(parents=True, exist_ok=True)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"[Export] Saved {len(items)} items to JSON: {filepath}")

def export_to_csv(items: List[ProductItem], filepath: Path) -> None:
    if not items:
        return
    filepath.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(items[0].model_dump().keys())
    with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        for item in items:
            writer.writerow(item.model_dump(mode="json"))
    print(f"[Export] Saved {len(items)} items to CSV: {filepath}")

def export_to_excel(items: List[ProductItem], filepath: Path) -> None:
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("[Export] openpyxl not installed, falling back to CSV.")
        csv_path = filepath.with_suffix(".csv")
        export_to_csv(items, csv_path)
        return

    filepath.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scraped Products"

    headers = ["Title", "Price (£)", "Currency", "Rating", "Availability", "Product URL", "Scraped At"]
    ws.append(headers)

    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
    center_align = Alignment(horizontal="center", vertical="center")

    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center_align

    for item in items:
        ws.append([
            item.title,
            item.price,
            item.currency,
            f"{item.rating} / 5",
            item.availability,
            str(item.product_url),
            item.scraped_at.strftime("%Y-%m-%d %H:%M:%S")
        ])

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

    wb.save(filepath)
    print(f"[Export] Styled Excel report saved: {filepath}")
